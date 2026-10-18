"""
Simulación de eventos discretos - Audifarma, sede Plaza del Sol
Jornada de 7:00 a 17:00, llegadas de Poisson con tasa variable por franja,
tiempos de servicio Gamma ajustados a la muestra de 420 atenciones,
cierre de un módulo de 12:00 a 14:00 y 200 réplicas independientes.

Requisitos:  pip install simpy numpy scipy matplotlib openpyxl
Uso:         python simulacion_audifarma.py
"""
import numpy as np
import simpy
import scipy.stats as st
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl

# ---------------------------------------------------------------- PARÁMETROS
SEMILLA = 2026
REPLICAS = 200
META_MIN = 10            # meta de servicio: espera promedio en cola (min)
UMBRAL_MIN = 15          # umbral para el % de usuarios que esperan "demasiado"
PROB_PREFERENCIAL = 0.08 # usuarios preferenciales (pasan primero)
TRANSICION_MIN = 0.0     # tiempo entre el llamado y el inicio de la atención (min)

FRANJAS = ["7:00", "8:00", "9:00", "10:00", "11:00",
           "12:00", "13:00", "14:00", "15:00", "16:00"]
LAMBDA = [89.8, 100.1, 89.1, 73.6, 55.1, 40.3, 48.7, 67.4, 57.3, 34.7]  # usuarios/h

# Módulos abiertos por franja (una lista por configuración a comparar)
CONFIGURACIONES = {
    "Actual": [4, 4, 4, 4, 4, 3, 3, 4, 4, 4],
    "Propuesta": [5, 6, 5, 4, 4, 3, 3, 4, 4, 3],   # Tabla 14 (meta Wq <= 10 min)
}

# ---------------------------------------------------- 1. AJUSTE DEL SERVICIO
x = np.loadtxt("tiempos_servicio.csv", skiprows=1)
candidatas = {
    "Exponencial": (st.expon, dict(floc=0)),
    "Gamma": (st.gamma, dict(floc=0)),
    "Lognormal": (st.lognorm, dict(floc=0)),
    "Weibull": (st.weibull_min, dict(floc=0)),
}
ajustes = []
for nombre, (dist, kw) in candidatas.items():
    p = dist.fit(x, **kw)
    ks = st.kstest(x, dist.cdf, args=p)
    bordes = dist.ppf(np.linspace(0, 1, 11), *p)          # 10 clases equiprobables
    obs = np.histogram(x, bordes)[0]
    esp = len(x) / 10
    chi2 = ((obs - esp) ** 2 / esp).sum()
    gl = 10 - 1 - (len(p) - 1)                              # la ubicación está fija en 0
    ajustes.append(dict(dist=nombre, params=p, D=ks.statistic, p_ks=ks.pvalue,
                        chi2=chi2, gl=gl, p_chi2=1 - st.chi2.cdf(chi2, gl)))
forma, _, escala = [a for a in ajustes if a["dist"] == "Gamma"][0]["params"]

# ---------------------------------------------------------- 2. EL MODELO
INICIO, FIN = 7 * 60, 17 * 60   # minutos desde medianoche


def una_jornada(modulos, rng):
    env = simpy.Environment(initial_time=INICIO)
    smax = max(modulos)
    sala = simpy.PriorityResource(env, capacity=smax)
    esperas, franja_lleg = [], []
    cola_min = {}                                            # tamaño de la fila minuto a minuto

    def usuario(franja):
        llegada = env.now
        prioridad = 0 if rng.random() < PROB_PREFERENCIAL else 1
        with sala.request(priority=prioridad) as turno:
            yield turno
            esperas.append(env.now - llegada)
            franja_lleg.append(franja)
            yield env.timeout(TRANSICION_MIN + rng.gamma(forma, escala))

    def cierre_modulo(desde, hasta, cuantos):
        # Un "bloqueo" con prioridad máxima toma el módulo cuando queda libre
        yield env.timeout(desde - env.now)
        with sala.request(priority=-1) as r:
            yield r
            yield env.timeout(max(0, hasta - env.now))

    def llegadas():
        t = INICIO
        while t < FIN:
            f = int((t - INICIO) // 60)
            tasa = LAMBDA[f] / 60                               # usuarios/min
            dt = rng.exponential(1 / tasa)
            limite = INICIO + (f + 1) * 60
            if t + dt >= limite:                                # propiedad sin memoria:
                t = limite                                      # se reinicia con la nueva tasa
                continue
            t += dt
            yield env.timeout(t - env.now)
            env.process(usuario(f))

    def monitor():
        while True:
            cola_min[int(env.now)] = sum(1 for r in sala.queue if r.priority >= 0)
            yield env.timeout(1)

    # Módulos cerrados: por cada franja con menos módulos que el máximo
    f = 0
    while f < 10:
        cerrados = smax - modulos[f]
        if cerrados > 0:
            g = f
            while g < 10 and smax - modulos[g] == cerrados:
                g += 1
            for _ in range(cerrados):
                env.process(cierre_modulo(INICIO + 60 * f, INICIO + 60 * g, 1))
            f = g
        else:
            f += 1

    env.process(llegadas())
    env.process(monitor())
    env.run(until=FIN + 240)                                    # deja vaciar la fila
    return np.array(esperas), np.array(franja_lleg), cola_min


# ------------------------------------------------------- 3. LAS 200 RÉPLICAS
rng = np.random.default_rng(SEMILLA)
resultados = {}
for nombre, modulos in CONFIGURACIONES.items():
    wq_dia, wq_franja, cola10, cola_max, todas, w_max, curvas = [], [], [], [], [], [], []
    minutos = np.arange(INICIO, FIN + 241)
    for _ in range(REPLICAS):
        w, fr, cola = una_jornada(modulos, rng)
        wq_dia.append(w.mean())
        wq_franja.append([w[fr == k].mean() if (fr == k).any() else 0 for k in range(10)])
        cola10.append(cola.get(10 * 60, 0))
        cola_max.append(max(cola.values()))
        w_max.append(w.max())
        curvas.append([cola.get(int(t), 0) for t in minutos])
        todas.append(w)
    todas = np.concatenate(todas)
    wq_dia = np.array(wq_dia)
    ic = 1.96 * wq_dia.std(ddof=1) / np.sqrt(REPLICAS)
    resultados[nombre] = dict(
        wq=wq_dia.mean(), ic=ic,
        wq_franja=np.array(wq_franja).mean(axis=0),
        pct_umbral=(todas > UMBRAL_MIN).mean(),
        p95=np.percentile(todas, 95),
        cola10=np.mean(cola10), cola_max=np.mean(cola_max),
        w_max=np.mean(w_max), curva=np.array(curvas).mean(axis=0), minutos=minutos,
        usuarios=len(todas) / REPLICAS, horas_modulo=sum(modulos))
    c = resultados[nombre]["curva"]
    resultados[nombre]["pico_curva"] = c.max()
    resultados[nombre]["hora_pico_curva"] = minutos[c.argmax()]
    resultados[nombre]["normaliza"] = next((t for t, v in zip(minutos, c) if t > minutos[c.argmax()] and v < 1), None)

# ------------------------------------------------------------ 4. RESULTADOS
print("AJUSTE DEL TIEMPO DE SERVICIO (n = %d)" % len(x))
for a in ajustes:
    print(f"  {a['dist']:12s} K-S D = {a['D']:.4f} (p = {a['p_ks']:.3f})   "
          f"chi2 = {a['chi2']:.2f}, gl = {a['gl']} (p = {a['p_chi2']:.3f})")
print(f"  -> Gamma: forma k = {forma:.3f}, escala θ = {escala:.4f} min "
      f"(media = {forma*escala:.2f} min)\n")
for nombre, r in resultados.items():
    print(f"CONFIGURACIÓN: {nombre}   ({REPLICAS} réplicas, {r['usuarios']:.0f} usuarios/día en promedio)")
    print(f"  Espera promedio en cola Wq : {r['wq']:.1f} min  (IC 95 %: ±{r['ic']:.1f})")
    print(f"  % que espera más de {UMBRAL_MIN} min : {100*r['pct_umbral']:.0f} %")
    print(f"  Percentil 95 de la espera  : {r['p95']:.1f} min")
    print(f"  Personas en fila a las 10:00: {r['cola10']:.0f}   (máximo del día: {r['cola_max']:.0f})")
    hp = r["hora_pico_curva"]; hn = r["normaliza"]
    print(f"  Curva promedio de la sala: máximo {r['pico_curva']:.1f} personas a las {hp//60}:{hp%60:02d}"
          + (f"; baja de 1 persona a las {hn//60}:{hn%60:02d}" if hn else ""))
    print(f"  Espera máxima del día (promedio de réplicas): {r['w_max']:.1f} min")
    print(f"  Horas-módulo por día: {r['horas_modulo']}")
    print("  Espera por franja de llegada:",
          ", ".join(f"{FRANJAS[k]}={v:.1f}" for k, v in enumerate(r["wq_franja"])), "\n")

# Excel con los resultados
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Ajuste_servicio"
ws.append(["Distribución", "Parámetros", "K-S D", "p (K-S)", "Chi-cuadrado", "gl", "p (chi2)"])
for a in ajustes:
    ws.append([a["dist"], ", ".join(f"{v:.4f}" for v in a["params"]),
               round(a["D"], 4), round(a["p_ks"], 4), round(a["chi2"], 2), a["gl"], round(a["p_chi2"], 4)])
ws2 = wb.create_sheet("Simulacion")
ws2.append(["Configuración", "Wq promedio (min)", "IC 95 % (±min)", f"% espera > {UMBRAL_MIN} min",
            "P95 espera (min)", "Fila a las 10:00", "Fila máxima"] + [f"Wq {f}" for f in FRANJAS])
for nombre, r in resultados.items():
    ws2.append([nombre, round(r["wq"], 2), round(r["ic"], 2), round(r["pct_umbral"], 3),
                round(r["p95"], 1), round(r["cola10"], 1), round(r["cola_max"], 1)]
               + [round(v, 1) for v in r["wq_franja"]])
if "Propuesta" in resultados:
    a, p = resultados["Actual"], resultados["Propuesta"]
    ws3 = wb.create_sheet("Tabla16")
    ws3.append(["Indicador", "Configuración actual", "Configuración propuesta", "Variación"])
    def var(x, y): return f"{(y / x - 1) * 100:+.0f} %"
    ws3.append(["Horas-módulo por día", a["horas_modulo"], p["horas_modulo"],
                f"{p['horas_modulo'] - a['horas_modulo']:+d}"])
    ws3.append(["Espera promedio en cola (min)", round(a["wq"], 1), round(p["wq"], 1), var(a["wq"], p["wq"])])
    ws3.append(["Percentil 95 de la espera (min)", round(a["p95"], 1), round(p["p95"], 1), var(a["p95"], p["p95"])])
    ws3.append([f"Usuarios que esperan más de {UMBRAL_MIN} min", f"{100*a['pct_umbral']:.1f} %",
                f"{100*p['pct_umbral']:.1f} %", f"{100*(p['pct_umbral']-a['pct_umbral']):+.1f} p.p."])
    ws3.append(["Espera máxima del día (min)", round(a["w_max"], 1), round(p["w_max"], 1), var(a["w_max"], p["w_max"])])
    ws3.append(["Máximo de usuarios en sala (curva promedio)", round(a["pico_curva"], 1), round(p["pico_curva"], 1),
                var(a["pico_curva"], p["pico_curva"])])
    ws3.append(["Máximo de usuarios en sala (promedio de los máximos diarios)", round(a["cola_max"], 1),
                round(p["cola_max"], 1), var(a["cola_max"], p["cola_max"])])
wb.save("resultados_simulacion.xlsx")

# Figura: espera promedio por franja (configuración actual)
r = resultados["Actual"]
fig, ax = plt.subplots(figsize=(8, 3.6), dpi=200)
barras = ax.bar(FRANJAS, r["wq_franja"], color="#d64545", width=0.8, zorder=2)
for b, v in zip(barras, r["wq_franja"]):
    ax.text(b.get_x() + b.get_width() / 2, v + 0.5, f"{v:.1f}".replace(".", ","), ha="center", va="bottom", fontsize=8)
ax.axhline(META_MIN, color="black", ls="--", lw=1, zorder=3)
ax.text(9.4, META_MIN + 0.6, f"Meta: {META_MIN} min", ha="right", fontsize=8)
ax.set_xlabel("Franja de llegada del usuario")
ax.set_ylabel("Espera promedio en fila (min)")
ax.grid(axis="y", color="#e5e5e5", zorder=0)
for s in ("top", "right"):
    ax.spines[s].set_visible(False)
fig.tight_layout()
fig.savefig("Figura5_espera_por_franja.png")

if "Propuesta" in resultados:
    a, p = resultados["Actual"], resultados["Propuesta"]
    ROJO, AZUL = "#d64545", "#1f3b73"

    # Figura 11: usuarios esperando en sala minuto a minuto
    fig, ax = plt.subplots(figsize=(8, 3.6), dpi=200)
    horas = a["minutos"] / 60
    lim = a["minutos"] <= 15 * 60
    ax.plot(horas[lim], a["curva"][lim], color=ROJO, lw=2, label="Configuración actual")
    ax.plot(horas[lim], p["curva"][lim], color=AZUL, lw=2, label="Configuración propuesta")
    for r_, c_ in ((a, ROJO), (p, AZUL)):
        hp = r_["hora_pico_curva"] / 60
        ax.annotate(f"{r_['pico_curva']:.1f}".replace(".", ","), (hp, r_["pico_curva"]),
                    xytext=(0, 6), textcoords="offset points", ha="center", fontsize=8, color=c_)
    ax.set_xticks(range(7, 16), [f"{h}:00" for h in range(7, 16)])
    ax.set_xlabel("Hora del día")
    ax.set_ylabel("Usuarios esperando en sala")
    ax.set_ylim(0, a["pico_curva"] * 1.15)
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#e5e5e5", zorder=0)
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    fig.tight_layout(); fig.savefig("Figura11_usuarios_en_sala.png")

    # Figura 12: espera por franja de llegada, actual vs propuesta
    fig, ax = plt.subplots(figsize=(8, 3.6), dpi=200)
    xs = np.arange(10); wdt = 0.4
    for off, r_, c_, lab in ((-wdt / 2, a, ROJO, "Configuración actual"), (wdt / 2, p, AZUL, "Configuración propuesta")):
        bs = ax.bar(xs + off, r_["wq_franja"], width=wdt - 0.03, color=c_, label=lab, zorder=2)
        for b, v in zip(bs, r_["wq_franja"]):
            ax.text(b.get_x() + b.get_width() / 2, v + 0.4, f"{v:.1f}".replace(".", ","), ha="center",
                    fontsize=7, zorder=4)
    ax.axhline(META_MIN, color="black", ls="--", lw=1, zorder=3)
    ax.text(9.6, META_MIN + 0.6, f"Meta: {META_MIN} min", ha="right", fontsize=8)
    ax.set_xticks(xs, FRANJAS)
    ax.set_xlabel("Franja de llegada del usuario")
    ax.set_ylabel("Espera promedio en cola (min)")
    ax.legend(frameon=False, loc="upper right", bbox_to_anchor=(1, 0.92))
    ax.grid(axis="y", color="#e5e5e5", zorder=0)
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    fig.tight_layout(); fig.savefig("Figura12_espera_actual_vs_propuesta.png")

print("Archivos: resultados_simulacion.xlsx, Figura5_espera_por_franja.png"
      + (", Figura11_usuarios_en_sala.png, Figura12_espera_actual_vs_propuesta.png" if "Propuesta" in resultados else ""))
