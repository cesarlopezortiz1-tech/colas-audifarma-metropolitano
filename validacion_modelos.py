"""
6.6 Comparación de los modelos
  - Figura 8: Wq de M/M/1, M/G/1 y M/M/s en los tres escenarios
  - Figura 9: Wq en la hora pico con 5 a 9 módulos
  - Tabla 12: validación de los modelos con simulación de larga duración en estado estacionario

Requisitos:  pip install numpy matplotlib openpyxl
Uso:         python validacion_modelos.py     (con tiempos_servicio.csv en la misma carpeta)
"""
import heapq
import math
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

x = np.loadtxt("tiempos_servicio.csv", skiprows=1)
media, sigma = x.mean(), x.std(ddof=1)          # min
mu = 60 / media                                  # usuarios/h
cs2 = (sigma / media) ** 2
forma = 1 / cs2; escala = media / forma          # Gamma con la media y la varianza observadas

ESCENARIOS = [("Hora valle", 34.7, 3), ("Demanda promedio", 65.6, 4), ("Hora pico", 100.1, 6)]
META = 10


# ------------------------------------------------------------ fórmulas (Wq en minutos)
def wq_mm1(lam, s):
    l = lam / s; r = l / mu
    return None if r >= 1 else r ** 2 / (1 - r) / l * 60

def wq_mg1(lam, s):
    l = lam / s; r = l / mu
    return None if r >= 1 else (l ** 2 * (sigma / 60) ** 2 + r ** 2) / (2 * (1 - r)) / l * 60

def wq_mms(lam, s):
    a = lam / mu; r = a / s
    if r >= 1: return None
    ult = a ** s / (math.factorial(s) * (1 - r))
    P0 = 1 / (sum(a ** n / math.factorial(n) for n in range(s)) + ult)
    return ult * P0 * r / (1 - r) / lam * 60

def wq_ac(lam, s):
    w = wq_mms(lam, s)
    return None if w is None else w * (1 + cs2) / 2


# ------------------------------------------------------------ simulación en estado estacionario
N_USUARIOS = 400_000      # usuarios por réplica
CALENTAMIENTO = 0.10      # fracción inicial descartada
REPLICAS = 10
rng = np.random.default_rng(2026)

def servicio(n, tipo):
    return rng.exponential(media, n) if tipo == "exp" else rng.gamma(forma, escala, n)

def sim_filas_independientes(lam, s, tipo):
    """Cada usuario se asigna al azar a un módulo: s colas M/G/1 con λ/s (recursión de Lindley)."""
    esperas = []
    for _ in range(s):
        n = N_USUARIOS // s
        a = rng.exponential(60 / (lam / s), n)           # tiempos entre llegadas (min)
        b = servicio(n, tipo)
        w = np.empty(n); w[0] = 0.0
        for i in range(1, n):
            w[i] = max(0.0, w[i - 1] + b[i - 1] - a[i])
        esperas.append(w[int(n * CALENTAMIENTO):])
    return np.concatenate(esperas).mean()

def sim_fila_unica(lam, s, tipo):
    """Fila única FIFO: cada usuario pasa al primer módulo que se libera."""
    n = N_USUARIOS
    t = np.cumsum(rng.exponential(60 / lam, n))
    b = servicio(n, tipo)
    libres = [0.0] * s; heapq.heapify(libres)
    w = np.empty(n)
    for i in range(n):
        f = heapq.heappop(libres)
        inicio = max(t[i], f)
        w[i] = inicio - t[i]
        heapq.heappush(libres, inicio + b[i])
    return w[int(n * CALENTAMIENTO):].mean()

def replicar(fun, *args):
    v = np.array([fun(*args) for _ in range(REPLICAS)])
    return v.mean(), 1.96 * v.std(ddof=1) / np.sqrt(REPLICAS)

prom, pico = ESCENARIOS[1], ESCENARIOS[2]
CASOS = [
    ("Filas independientes, servicio exponencial", prom, sim_filas_independientes, "exp", [("M/M/1", wq_mm1)]),
    ("Filas independientes, servicio gamma", prom, sim_filas_independientes, "gamma", [("M/G/1", wq_mg1)]),
    ("Filas independientes, servicio gamma", pico, sim_filas_independientes, "gamma", [("M/G/1", wq_mg1)]),
    ("Fila única, servicio exponencial", prom, sim_fila_unica, "exp", [("M/M/s", wq_mms)]),
    ("Fila única, servicio gamma", prom, sim_fila_unica, "gamma", [("M/M/s", wq_mms), ("A-C", wq_ac)]),
    ("Fila única, servicio gamma", pico, sim_fila_unica, "gamma", [("M/M/s", wq_mms), ("A-C", wq_ac)]),
]

print(f"Servicio: media = {media:.4f} min, σ = {sigma:.4f} min, μ = {mu:.4f}/h, Cs² = {cs2:.4f}\n")
print("TABLA 12. VALIDACIÓN POR SIMULACIÓN (Wq en minutos)")
tabla12 = []
for conf, (esc, lam, s), fun, tipo, modelos in CASOS:
    sim, ic = replicar(fun, lam, s, tipo)
    analit = " | ".join(f"{m}: {f(lam, s):.2f}" for m, f in modelos)
    ref = modelos[-1][1](lam, s)
    dif = (sim - ref) / ref * 100
    tabla12.append([conf, f"{esc}, s = {s}", analit, round(sim, 2), round(ic, 2), f"{dif:+.1f} % (vs. {modelos[-1][0]})"])
    print(f"  {conf} | {esc}, s = {s} | {analit} | simulado {sim:.2f} ± {ic:.2f} | {dif:+.1f} % vs {modelos[-1][0]}")

# ------------------------------------------------------------ módulos en hora pico
print("\nESPERA EN HORA PICO (λ = 100,1) SEGÚN MÓDULOS")
mods = list(range(5, 10))
curvas = {"M/M/1": [], "M/G/1": [], "M/M/s": [], "M/M/s Allen-Cunneen": []}
for s in mods:
    for k, f in [("M/M/1", wq_mm1), ("M/G/1", wq_mg1), ("M/M/s", wq_mms), ("M/M/s Allen-Cunneen", wq_ac)]:
        curvas[k].append(f(100.1, s))
    print(f"  s = {s}: " + ", ".join(f"{k} = {'inestable' if v[-1] is None else f'{v[-1]:.2f}'}"
                                     for k, v in curvas.items()))
for k, v in curvas.items():
    s_min = next((s for s, w in zip(mods, v) if w is not None and w <= META), None)
    print(f"  Módulos mínimos para Wq ≤ {META} min con {k}: {s_min}")

# ------------------------------------------------------------ Excel
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Tabla12"
ws.append(["Configuración simulada", "Escenario", "Resultado analítico", "Resultado simulado",
           "IC 95 % (±)", "Diferencia"])
for f in tabla12: ws.append(f)
ws2 = wb.create_sheet("Modulos_hora_pico")
ws2.append(["Módulos"] + list(curvas))
for i, s in enumerate(mods):
    ws2.append([s] + [("Inestable" if v[i] is None else round(v[i], 2)) for v in curvas.values()])
ws3 = wb.create_sheet("Figura8")
ws3.append(["Escenario", "M/M/1", "M/G/1", "M/M/s"])
for e, lam, s in ESCENARIOS:
    ws3.append([f"{e} (s = {s})", round(wq_mm1(lam, s), 2), round(wq_mg1(lam, s), 2), round(wq_mms(lam, s), 2)])
wb.save("resultados_validacion.xlsx")

# ------------------------------------------------------------ Figura 8
colores = {"M/M/1": "#d64545", "M/G/1": "#e8a23a", "M/M/s": "#1f3b73"}
fig, ax = plt.subplots(figsize=(8, 3.8), dpi=200)
xs = np.arange(len(ESCENARIOS)); w = 0.26
for j, (k, f) in enumerate([("M/M/1", wq_mm1), ("M/G/1", wq_mg1), ("M/M/s", wq_mms)]):
    vals = [f(l, s) for _, l, s in ESCENARIOS]
    b = ax.bar(xs + (j - 1) * w, vals, width=w - 0.02, color=colores[k], label=k, zorder=2)
    for r, v in zip(b, vals):
        ax.text(r.get_x() + r.get_width() / 2, v + 0.3, f"{v:.1f}".replace(".", ","), ha="center", fontsize=8,
                    zorder=4, bbox=dict(facecolor="white", edgecolor="none", pad=1))
ax.axhline(META, color="black", ls="--", lw=1, zorder=3)
ax.text(2.45, META + 0.4, f"Meta: {META} min", ha="right", fontsize=8, zorder=4)
ax.set_xticks(xs, [f"{e}\n(λ = {str(l).replace('.', ',')}; s = {s})" for e, l, s in ESCENARIOS])
ax.set_ylabel("Espera promedio en cola (min)")
ax.legend(frameon=False, loc="upper left")
ax.grid(axis="y", color="#e5e5e5", zorder=0)
for s_ in ("top", "right"): ax.spines[s_].set_visible(False)
fig.tight_layout(); fig.savefig("Figura8_comparacion_modelos.png")

# ------------------------------------------------------------ Figura 9
fig, ax = plt.subplots(figsize=(8, 3.8), dpi=200)
for k, c, ls in [("M/M/1", "#d64545", "-"), ("M/G/1", "#e8a23a", "-"), ("M/M/s", "#1f3b73", "-")]:
    ss = [s for s, v in zip(mods, curvas[k]) if v is not None]
    vv = [v for v in curvas[k] if v is not None]
    ax.plot(ss, vv, color=c, ls=ls, lw=2, marker="o", ms=6, label=k, zorder=3)
    for s, v in zip(ss, vv):
        ax.text(s + 0.06, v + 0.4, f"{v:.1f}".replace(".", ","), fontsize=8, color="#333333")
ax.axhline(META, color="black", ls="--", lw=1, zorder=2)
ax.text(9.0, META + 0.4, f"Meta: {META} min", ha="right", fontsize=8)
ax.set_xticks(mods, [f"{s}" for s in mods])
ax.set_xlabel("Módulos habilitados en la hora pico (λ = 100,1 usuarios/h)")
ax.set_ylabel("Espera promedio en cola (min)")
ax.text(5, 1, "Con 5 módulos\nel sistema es inestable", fontsize=8, ha="center", va="bottom", color="#555555")
ax.set_xlim(4.5, 9.5)
ax.legend(frameon=False)
ax.grid(axis="y", color="#e5e5e5", zorder=0)
for s_ in ("top", "right"): ax.spines[s_].set_visible(False)
fig.tight_layout(); fig.savefig("Figura9_modulos_hora_pico.png")
print("\nArchivos: resultados_validacion.xlsx, Figura8_comparacion_modelos.png, Figura9_modulos_hora_pico.png")
