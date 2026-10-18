"""
6.2 Ajuste de distribuciones - tiempos de servicio + Tabla 8 (resumen de todas las pruebas)
  - Exponencial y Gamma sobre los 420 tiempos de servicio (K-S)
  - Valores críticos al 5 % por Monte Carlo, porque los parámetros se estiman con los mismos datos
  - Figura 7: histograma con las dos curvas ajustadas
  - Tabla 8: resumen de llegadas (Poisson, exponencial) y servicio (exponencial, gamma)

Requisitos:  pip install numpy scipy matplotlib openpyxl
Uso:         python prueba_servicio.py
             (con tiempos_servicio.csv y Datos_Audifarma.xlsx en la misma carpeta)
"""
import datetime as dt
import numpy as np
import scipy.stats as st
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

rng = np.random.default_rng(1)
MC = 2000   # muestras Monte Carlo para los valores críticos

# ------------------------------------------------------------ datos
x = np.loadtxt("tiempos_servicio.csv", skiprows=1)
wb = openpyxl.load_workbook("Datos_Audifarma.xlsx")
iv = np.array([r[1] for r in wb["Intervalos_5min"].iter_rows(min_row=4, values_only=True)
               if isinstance(r[1], (int, float))], int)

def seg(t):
    return t.hour * 3600 + t.minute * 60 + t.second + t.microsecond / 1e6

ia = np.array([(seg(r[2]) - seg(r[3])) / 60
               for r in wb["Entre_llegadas"].iter_rows(min_row=4, values_only=True)
               if isinstance(r[2], dt.time)])


def ks_exponencial(datos):
    D = st.kstest(datos, "expon", args=(0, datos.mean())).statistic
    sim = np.array([(lambda y: st.kstest(y, "expon", args=(0, y.mean())).statistic)
                    (rng.exponential(1, len(datos))) for _ in range(MC)])
    return D, np.percentile(sim, 95), (sim >= D).mean()


def ks_gamma(datos):
    a, _, th = st.gamma.fit(datos, floc=0)
    D = st.kstest(datos, "gamma", args=(a, 0, th)).statistic
    sim = []
    for _ in range(MC):
        y = rng.gamma(a, th, len(datos))
        aa, _, tt = st.gamma.fit(y, floc=0)
        sim.append(st.kstest(y, "gamma", args=(aa, 0, tt)).statistic)
    sim = np.array(sim)
    return a, th, D, np.percentile(sim, 95), (sim >= D).mean()


# ------------------------------------------------------------ pruebas
# Servicio
D_exp, c_exp, p_exp = ks_exponencial(x)
alfa, theta, D_gam, c_gam, p_gam = ks_gamma(x)

# Llegadas: chi-cuadrado Poisson (mismas clases que la Tabla 7)
n, m = len(iv), iv.mean()
lo = next(k for k in range(30) if n * st.poisson.cdf(k, m) >= 5)
hi = next(k for k in range(40, 0, -1) if n * st.poisson.sf(k - 1, m) >= 5)
obs = np.array([(iv <= lo).sum()] + [(iv == k).sum() for k in range(lo + 1, hi)] + [(iv >= hi).sum()])
esp = n * np.array([st.poisson.cdf(lo, m)] + [st.poisson.pmf(k, m) for k in range(lo + 1, hi)]
                   + [st.poisson.sf(hi - 1, m)])
chi2 = ((obs - esp) ** 2 / esp).sum()
gl = len(obs) - 2
c_chi = st.chi2.ppf(0.95, gl)

# Llegadas: tiempos entre llegadas
D_ia, c_ia, p_ia = ks_exponencial(ia)

# ------------------------------------------------------------ resultados
print("TIEMPOS DE SERVICIO (n = %d, media = %.3f min, desv. = %.3f min)" % (len(x), x.mean(), x.std(ddof=1)))
print(f"  Exponencial (media {x.mean():.2f}): D = {D_exp:.4f}, crítico = {c_exp:.4f}, p = {p_exp:.3f}"
      f"  -> {'Se rechaza' if D_exp > c_exp else 'No se rechaza'}")
print(f"  Gamma (α = {alfa:.2f}, θ = {theta:.2f}): D = {D_gam:.4f}, crítico = {c_gam:.4f}, p = {p_gam:.3f}"
      f"  -> {'Se rechaza' if D_gam > c_gam else 'No se rechaza'}")
print("  (valores críticos y p por Monte Carlo, con parámetros estimados de los datos)\n")

filas = [
    ["Llegadas por intervalo de 5 min", f"Poisson (λ = {m+1e-9:.2f})", "Chi-cuadrado",
     round(chi2, 2), round(c_chi, 2), "No se rechaza" if chi2 < c_chi else "Se rechaza"],
    ["Tiempo entre llegadas", f"Exponencial (media {ia.mean():.3f} min)", "K-S",
     round(D_ia, 4), round(c_ia, 4), "No se rechaza" if D_ia < c_ia else "Se rechaza"],
    ["Tiempo de servicio", f"Exponencial (media {x.mean():.2f} min)", "K-S",
     round(D_exp, 4), round(c_exp, 4), "No se rechaza" if D_exp < c_exp else "Se rechaza"],
    ["Tiempo de servicio", f"Gamma (α = {alfa:.2f}; θ = {theta:.2f})", "K-S",
     round(D_gam, 4), round(c_gam, 4), "No se rechaza" if D_gam < c_gam else "Se rechaza"],
]
cab = ["Variable", "Distribución probada", "Prueba", "Estadístico", "Valor crítico", "Decisión"]
print("TABLA 8. RESUMEN DE LAS PRUEBAS DE BONDAD DE AJUSTE (α = 0,05)")
for f in filas:
    print("  " + " | ".join(str(v) for v in f))

out = openpyxl.Workbook()
ws = out.active; ws.title = "Tabla8"
ws.append(cab)
for f in filas:
    ws.append(f)
out.save("resultados_servicio.xlsx")

# ------------------------------------------------------------ Figura 7
fig, ax = plt.subplots(figsize=(8, 3.8), dpi=200)
ax.hist(x, bins=np.arange(0, x.max() + 0.5, 0.5), density=True, color="#d64545", alpha=0.75,
        edgecolor="white", label="Datos observados", zorder=2)
t = np.linspace(0.001, x.max() + 0.5, 400)
ax.plot(t, st.expon.pdf(t, 0, x.mean()), color="#555555", lw=2, ls="--",
        label=f"Exponencial (media {x.mean():.2f} min)".replace(".", ","), zorder=3)
ax.plot(t, st.gamma.pdf(t, alfa, 0, theta), color="#1f3b73", lw=2,
        label=f"Gamma (α = {alfa:.2f}; θ = {theta:.2f})".replace(".", ","), zorder=3)
ax.set_xlabel("Tiempo de servicio (min)")
ax.set_ylabel("Densidad")
ax.legend(frameon=False)
ax.grid(axis="y", color="#e5e5e5", zorder=0)
for s in ("top", "right"):
    ax.spines[s].set_visible(False)
fig.tight_layout()
fig.savefig("Figura7_tiempos_servicio.png")
print("\nArchivos: resultados_servicio.xlsx, Figura7_tiempos_servicio.png")
