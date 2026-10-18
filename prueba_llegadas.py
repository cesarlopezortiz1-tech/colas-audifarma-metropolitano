"""
6.2 Ajuste de distribuciones - proceso de llegadas (franja 9:00-10:00, 10 días)
  1) Prueba chi-cuadrado de Poisson sobre las llegadas por intervalo de 5 min
  2) Prueba K-S (corrección de Lilliefors) de la exponencial sobre los tiempos entre llegadas

Requisitos:  pip install numpy scipy matplotlib openpyxl
Uso:         python prueba_llegadas.py      (con Datos_Audifarma.xlsx en la misma carpeta)
"""
import datetime as dt
import numpy as np
import scipy.stats as st
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

wb = openpyxl.load_workbook("Datos_Audifarma.xlsx")

# ------------------------------------------------ datos
iv = np.array([r[1] for r in wb["Intervalos_5min"].iter_rows(min_row=4, values_only=True)
               if isinstance(r[1], (int, float))], int)

def seg(t):
    return t.hour * 3600 + t.minute * 60 + t.second + t.microsecond / 1e6

ia = np.array([(seg(r[2]) - seg(r[3])) / 60
               for r in wb["Entre_llegadas"].iter_rows(min_row=4, values_only=True)
               if isinstance(r[2], dt.time)])

# ------------------------------------------------ 1) chi-cuadrado Poisson
n, m, v = len(iv), iv.mean(), iv.var(ddof=1)
lo = next(k for k in range(30) if n * st.poisson.cdf(k, m) >= 5)          # agrupa la cola izquierda
hi = next(k for k in range(40, 0, -1) if n * st.poisson.sf(k - 1, m) >= 5)  # agrupa la cola derecha
etiquetas = [f"≤{lo}"] + [str(k) for k in range(lo + 1, hi)] + [f"≥{hi}"]
obs = np.array([(iv <= lo).sum()] + [(iv == k).sum() for k in range(lo + 1, hi)] + [(iv >= hi).sum()])
esp = n * np.array([st.poisson.cdf(lo, m)] + [st.poisson.pmf(k, m) for k in range(lo + 1, hi)]
                   + [st.poisson.sf(hi - 1, m)])
chi2 = ((obs - esp) ** 2 / esp).sum()
gl = len(obs) - 1 - 1                    # se estimó un parámetro (la media)
crit = st.chi2.ppf(0.95, gl)
p_chi = st.chi2.sf(chi2, gl)

# ------------------------------------------------ 2) K-S / Lilliefors exponencial
D = st.kstest(ia, "expon", args=(0, ia.mean())).statistic
rng = np.random.default_rng(1)
sim = np.array([st.kstest(y, "expon", args=(0, y.mean())).statistic
                for y in (rng.exponential(1, len(ia)) for _ in range(10000))])
D_crit = np.percentile(sim, 95)          # valor crítico de Lilliefors al 5 % (Monte Carlo)
p_ks = (sim >= D).mean()

# ------------------------------------------------ resultados
print("PRUEBA CHI-CUADRADO - LLEGADAS POR INTERVALO DE 5 MIN (9:00-10:00)")
print(f"  Intervalos: {n}   Media: {m:.3f}   Varianza: {v:.3f}   Varianza/media: {v/m:.3f}")
print(f"  {'Llegadas':>9} {'Observada':>10} {'Esperada':>10}")
for e, o, x in zip(etiquetas, obs, esp):
    print(f"  {e:>9} {o:>10d} {x:>10.2f}")
print(f"  {'Total':>9} {obs.sum():>10d} {esp.sum():>10.2f}")
print(f"  Chi-cuadrado = {chi2:.2f}, gl = {gl}, valor crítico (5 %) = {crit:.2f}, p = {p_chi:.3f}")
print("  -> No se rechaza Poisson" if chi2 < crit else "  -> Se rechaza Poisson")
print()
print("PRUEBA K-S (LILLIEFORS) - TIEMPOS ENTRE LLEGADAS (9:00-10:00)")
print(f"  n = {len(ia)}   Media = {ia.mean():.3f} min   Desv. est. = {ia.std(ddof=1):.3f} min   "
      f"CV = {ia.std(ddof=1)/ia.mean():.3f}")
print(f"  D = {D:.4f}, valor crítico de Lilliefors (5 %) = {D_crit:.4f}, p = {p_ks:.3f}")
print("  -> No se rechaza la exponencial" if D < D_crit else "  -> Se rechaza la exponencial")

# Tabla 7 en Excel
out = openpyxl.Workbook()
ws = out.active; ws.title = "Tabla7_chi2_Poisson"
ws.append(["Llegadas por intervalo", "Frecuencia observada", "Frecuencia esperada"])
for e, o, x in zip(etiquetas, obs, esp):
    ws.append([e, int(o), round(float(x), 2)])
ws.append(["Total", int(obs.sum()), round(float(esp.sum()), 2)])
ws.append([])
for lab, val in [("Media", m), ("Varianza", v), ("Chi-cuadrado", chi2), ("Grados de libertad", gl),
                 ("Valor crítico 5 %", crit), ("Valor p", p_chi)]:
    ws.append([lab, round(float(val), 4)])
ws2 = out.create_sheet("KS_entre_llegadas")
for lab, val in [("n", len(ia)), ("Media (min)", ia.mean()), ("Desv. est. (min)", ia.std(ddof=1)),
                 ("D", D), ("Valor crítico Lilliefors 5 %", D_crit), ("Valor p", p_ks)]:
    ws2.append([lab, round(float(val), 4)])
out.save("resultados_llegadas.xlsx")

# Figura 6
x = np.arange(len(obs))
fig, ax = plt.subplots(figsize=(8, 3.6), dpi=200)
ax.bar(x - 0.2, obs, width=0.4, color="#d64545", label="Observada", zorder=2)
ax.bar(x + 0.2, esp, width=0.4, color="#9aa5b1", label="Esperada (Poisson)", zorder=2)
ax.set_xticks(x, etiquetas)
ax.set_xlabel("Llegadas por intervalo de 5 minutos")
ax.set_ylabel("Frecuencia (intervalos)")
ax.legend(frameon=False)
ax.grid(axis="y", color="#e5e5e5", zorder=0)
for s in ("top", "right"):
    ax.spines[s].set_visible(False)
fig.tight_layout()
fig.savefig("Figura6_llegadas_poisson.png")
print("\nArchivos: resultados_llegadas.xlsx, Figura6_llegadas_poisson.png")
