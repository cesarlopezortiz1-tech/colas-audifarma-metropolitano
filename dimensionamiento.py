"""
6.7 Dimensionamiento de módulos por franja y análisis económico
  - Tabla 14: módulos mínimos por franja para Wq <= 10 min (M/M/s y Allen-Cunneen)
  - Tabla 15: costo total por hora en la franja pico según el número de módulos
  - Figura 10: costo de servicio, de espera y total en la franja pico
  - Extra: óptimo económico por franja

Requisitos:  pip install numpy matplotlib openpyxl
Uso:         python dimensionamiento.py     (con tiempos_servicio.csv en la misma carpeta)
"""
import math
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

x = np.loadtxt("tiempos_servicio.csv", skiprows=1)
media, sigma = x.mean(), x.std(ddof=1)
mu = 60 / media
factor_ac = (1 + (sigma / media) ** 2) / 2

META = 10            # min
S_MAX = 6            # módulos físicos
C_SERV = 19_000      # COP por hora-módulo (supuesto)
C_ESP = 9_620        # COP por hora de espera de un usuario (SMMLV 2026 / 182 h)

FRANJAS = ["7:00-8:00", "8:00-9:00", "9:00-10:00", "10:00-11:00", "11:00-12:00",
           "12:00-13:00", "13:00-14:00", "14:00-15:00", "15:00-16:00", "16:00-17:00"]
LAMBDA = [89.8, 100.1, 89.1, 73.6, 55.1, 40.3, 48.7, 67.4, 57.3, 34.7]
S_ACT = [4, 4, 4, 4, 4, 3, 3, 4, 4, 4]


def mms(lam, s):
    """Devuelve (ρ, Lq, Wq en min) del M/M/s, o None si es inestable."""
    a = lam / mu; r = a / s
    if r >= 1: return None
    ult = a ** s / (math.factorial(s) * (1 - r))
    P0 = 1 / (sum(a ** n / math.factorial(n) for n in range(s)) + ult)
    lq = ult * P0 * r / (1 - r)
    return r, lq, lq / lam * 60

def costo(lam, s):
    m = mms(lam, s)
    return None if m is None else (C_SERV * s, m[1], m[1] * C_ESP, C_SERV * s + m[1] * C_ESP)

def fmt(v, d=1):
    return f"{v:.{d}f}".replace(".", ",")

wb = openpyxl.Workbook()

# ------------------------------------------------------------ Tabla 14
ws = wb.active; ws.title = "Tabla14"
cab = ["Franja", "λ (usuarios/h)", "s actual", "Wq actual M/M/s (min)", "s propuesto",
       "ρ propuesto", "Wq M/M/s (min)", "Wq M/G/s A-C (min)"]
ws.append(cab)
print("TABLA 14. MÓDULOS PROPUESTOS POR FRANJA (meta Wq ≤ %d min, máximo %d módulos)" % (META, S_MAX))
print("  " + " | ".join(cab))
s_prop = []
for f, lam, sa in zip(FRANJAS, LAMBDA, S_ACT):
    act = mms(lam, sa)
    sp = next(s for s in range(1, S_MAX + 1) if mms(lam, s) and mms(lam, s)[2] <= META)
    r, lq, wq = mms(lam, sp)
    s_prop.append(sp)
    fila = [f, lam, sa, "Inestable" if act is None else round(act[2], 2), sp, round(r, 3), round(wq, 2),
            round(wq * factor_ac, 2)]
    ws.append(fila)
    print("  " + " | ".join(fmt(v, 3 if i == 5 else 1) if isinstance(v, float) and i != 1 else str(v).replace(".", ",")
                          for i, v in enumerate(fila)))
ws.append(["Horas-módulo / día", "", sum(S_ACT), "", sum(s_prop)])
print(f"  Horas-módulo / día: actual {sum(S_ACT)}, propuesto {sum(s_prop)}\n")

# ------------------------------------------------------------ Tabla 15 (hora pico) y demanda promedio
ws2 = wb.create_sheet("Tabla15")
ws2.append(["Módulos (s)", "Costo de servicio (COP/h)", "Lq (usuarios)", "Costo de espera (COP/h)",
            "Costo total (COP/h)"])
print("TABLA 15. COSTO TOTAL POR HORA EN LA FRANJA PICO (λ = 100,1)")
pico = {}
for s in range(6, 10):
    cs, lq, ce, ct = costo(100.1, s)
    pico[s] = (cs, ce, ct)
    ws2.append([s, cs, round(lq, 2), round(ce), round(ct)])
    print(f"  s = {s}: servicio ${cs:,.0f} | Lq {lq:.2f} | espera ${ce:,.0f} | total ${ct:,.0f}".replace(",", "."))
s_opt = min(pico, key=lambda s: pico[s][2])
print(f"  Óptimo: {s_opt} módulos; diferencia de 6 frente al óptimo: "
      f"{(pico[6][2] / pico[s_opt][2] - 1) * 100:.1f} %\n")

print("DEMANDA PROMEDIO (λ = 65,6)")
for s in range(4, 8):
    cs, lq, ce, ct = costo(65.6, s)
    print(f"  s = {s}: total ${ct:,.0f}".replace(",", "."))

# ------------------------------------------------------------ óptimo económico por franja
ws3 = wb.create_sheet("Optimo_por_franja")
ws3.append(["Franja", "λ", "s propuesto (meta 10 min)", "Costo con s propuesto", "s de mínimo costo",
            "Costo mínimo"])
print("\nÓPTIMO ECONÓMICO POR FRANJA")
for f, lam, sp in zip(FRANJAS, LAMBDA, s_prop):
    c = {s: costo(lam, s)[3] for s in range(1, 10) if costo(lam, s)}
    so = min(c, key=c.get)
    ws3.append([f, lam, sp, round(c[sp]), so, round(c[so])])
    print(f"  {f}: propuesto {sp} (${c[sp]:,.0f}) | mínimo costo {so} (${c[so]:,.0f})".replace(",", "."))
wb.save("resultados_dimensionamiento.xlsx")

# ------------------------------------------------------------ Figura 10
ss = list(pico)
fig, ax = plt.subplots(figsize=(8, 3.8), dpi=200)
ax.plot(ss, [pico[s][0] / 1000 for s in ss], color="#1f3b73", lw=2, marker="o", label="Costo de servicio")
ax.plot(ss, [pico[s][1] / 1000 for s in ss], color="#e8a23a", lw=2, marker="o", label="Costo de espera")
ax.plot(ss, [pico[s][2] / 1000 for s in ss], color="#d64545", lw=2.5, marker="o", label="Costo total")
for s in ss:
    ax.text(s, pico[s][2] / 1000 + 5, f"${pico[s][2]/1000:,.1f} mil".replace(".", ","), ha="center",
            fontsize=8, color="#333333")
ax.annotate("Mínimo", (s_opt, pico[s_opt][2] / 1000), xytext=(s_opt + 0.35, pico[s_opt][2] / 1000 - 30),
            arrowprops=dict(arrowstyle="->", color="#555555"), fontsize=8, color="#555555")
ax.set_xticks(ss)
ax.set_xlabel("Módulos habilitados en la franja pico (8:00-9:00)")
ax.set_ylabel("Miles de pesos por hora")
ax.set_ylim(0, max(pico[s][2] for s in ss) / 1000 * 1.15)
ax.legend(frameon=False, loc="center right")
ax.grid(axis="y", color="#e5e5e5", zorder=0)
for s_ in ("top", "right"): ax.spines[s_].set_visible(False)
fig.tight_layout(); fig.savefig("Figura10_costos_hora_pico.png")
print("\nArchivos: resultados_dimensionamiento.xlsx, Figura10_costos_hora_pico.png")
